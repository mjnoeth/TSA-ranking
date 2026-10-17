#!/usr/bin/env python3
"""Extract next-season team rankings from a legacy Dual Meet Summary workbook.

Reads the "North Report" and "South Report" sheets of the previous season's
workbook and writes north_ranks.csv / south_ranks.csv with columns rank,team.

The Report sheets have a fixed layout across all known years:
  row 1-3  headers
  row 4+   col A = proposed next-season rank, col B = team name
Rows whose team name starts with "**" are placeholders and are skipped.

Usage:
  python src/tools/legacy_extract_rank.py <workbook.xlsx> [--out-dir DIR]

By default the CSVs are written next to the workbook.
"""

from __future__ import annotations

import argparse
import csv
import sys
import warnings
from pathlib import Path

import openpyxl

SHEETS = {"North Report": "north_ranks.csv", "South Report": "south_ranks.csv"}
FIRST_DATA_ROW = 4


def extract(ws) -> list[tuple[int, str]]:
    teams: list[tuple[int, str]] = []
    for row in ws.iter_rows(min_row=FIRST_DATA_ROW, max_col=2, values_only=True):
        rank, team = row
        if not isinstance(team, str) or not team.strip():
            continue
        team = team.strip()
        if team.startswith("**"):
            continue
        if team.startswith("#"):
            raise ValueError(f"{ws.title}: cell error {team!r} in team column; workbook is incomplete")
        if not isinstance(rank, (int, float)):
            raise ValueError(f"{ws.title}: non-numeric rank {rank!r} for team {team!r}")
        teams.append((int(rank), team))

    ranks = [r for r, _ in teams]
    if ranks != list(range(1, len(ranks) + 1)):
        raise ValueError(f"{ws.title}: ranks are not contiguous 1..N: {ranks}")
    names = [t for _, t in teams]
    if len(set(names)) != len(names):
        dupes = sorted({n for n in names if names.count(n) > 1})
        raise ValueError(f"{ws.title}: duplicate team names: {dupes}")
    return teams


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("workbook", type=Path, help="previous season's Dual Meet Summary .xlsx")
    p.add_argument("--out-dir", type=Path, help="directory for CSV output (default: workbook's directory)")
    args = p.parse_args(argv)

    if not args.workbook.is_file():
        p.error(f"workbook not found: {args.workbook}")
    out_dir = args.out_dir or args.workbook.parent
    out_dir.mkdir(parents=True, exist_ok=True)

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # openpyxl print-area noise
        wb = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)

    missing = [s for s in SHEETS if s not in wb.sheetnames]
    if missing:
        print(f"error: workbook is missing required sheet(s): {missing}", file=sys.stderr)
        print(f"       sheets present: {wb.sheetnames}", file=sys.stderr)
        return 1

    for sheet, filename in SHEETS.items():
        try:
            teams = extract(wb[sheet])
        except ValueError as e:
            print(f"error: {e}", file=sys.stderr)
            return 1
        out = out_dir / filename
        with out.open("w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(["rank", "team"])
            w.writerows(teams)
        print(f"{sheet}: {len(teams)} teams -> {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
